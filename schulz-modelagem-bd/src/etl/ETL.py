#Instalação das dependências
!pip install -q gradio pandas openpyxl

import gradio as gr
import pandas as pd
import openpyxl
import re
import datetime

# Pipeline ETL da SCHULZ S.A.

def pipeline_etl_schulz(arquivo):
    if arquivo is None:
        return "Nenhum arquivo enviado.", None, None, None

    caminho_arquivo = arquivo.name
    nome_arquivo_origem = caminho_arquivo.split("/")[-1]

    log_execucao = []
    log_execucao.append(f"INÍCIO DO PROCESSAMENTO ETL")
    log_execucao.append(f"• Preservando arquivo original: '{nome_arquivo_origem}' (Read-Only)")

    # 1. ETAPA DE EXTRAÇÃO (EXTRACT) COM MAPEAMENTO DE ORIGEM
    try:
        wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)
        nome_aba = wb.sheetnames[0]
        sheet = wb[nome_aba]
    except Exception as e:
        return f"Erro ao abrir o arquivo: {str(e)}", None, None, None

    registros_extraidos = []
    pendencias_registradas = []

    # Identificação do arquivo: se é relatório MMC ou Formulário de Dispositivo
    is_mmc = "Calypso" in str(sheet['B2'].value or '') or "Measurement Plan" in str(sheet['B3'].value or '')

    if is_mmc:
        log_execucao.append("• Fonte identificada: Relatório da Máquina de Medição por Coordenadas (MMC)")

        # Leitura dos Metadados do Cabeçalho MMC
        plano_medicao = str(sheet['B4'].value or '').strip()
        data_raw = sheet['D4'].value
        metrologista = str(sheet['B10'].value or '').strip()
        cod_mmc = str(sheet['D10'].value or '').strip() # Mantido como string/código
        fornecedor = str(sheet['F10'].value or '').strip()
        ordem_servico = str(sheet['F4'].value or '').strip()
        part_no = str(sheet['F7'].value or '').strip()

        # Varredura da tabela de medições a partir da linha 13
        for r in range(13, sheet.max_row + 1):
            char_name = sheet.cell(row=r, column=1).value
            if not char_name:
                continue

            registro = {
                "FONTE_ORIGEM": f"{nome_arquivo_origem} -> Aba: {nome_aba}",
                "PLANO_MEDICAO": plano_medicao,
                "METROLOGISTA": metrologista,
                "COD_MMC": str(cod_mmc), # Código
                "FORNECEDOR": fornecedor,
                "ORDEM_SERVICO": str(ordem_servico), # Código
                "PART_NO": str(part_no), # Código
                "CARACTERISTICA": str(char_name).strip(),
                "VALOR_OBTIDO": sheet.cell(row=r, column=2).value,
                "VALOR_NOMINAL": sheet.cell(row=r, column=3).value,
                "TOLERANCIA_SUP": sheet.cell(row=r, column=4).value,
                "TOLERANCIA_INF": sheet.cell(row=r, column=5).value,
                "DESVIO": sheet.cell(row=r, column=6).value,
                "DATA_MEDICAO": data_raw
            }
            registros_extraidos.append(registro)

    else:
        log_execucao.append("• Fonte identificada: Formulário de Calibração / Dispositivo de Controle")
        # Processamento flexível para estruturas desestruturadas
        for r in range(1, min(sheet.max_row + 1, 100)):
            linha_vals = [sheet.cell(row=r, column=c).value for c in range(1, min(sheet.max_column + 1, 15))]
            if any(v is not None for v in linha_vals):
                registros_extraidos.append({
                    "FONTE_ORIGEM": f"{nome_arquivo_origem} -> Linha {r}",
                    "DADOS_BRUTOS": " | ".join([str(v) for v in linha_vals if v is not None])
                })

    df_bruto = pd.DataFrame(registros_extraidos)

    # 2. ETAPA DE TRANSFORMAÇÃO DOS DADOS (TRANSFORM)
    log_execucao.append("\n ETAPA DE TRANSFORMAÇÃO E HIGIENIZAÇÃO ")
    df_sanitizado = df_bruto.copy()

    # Função de limpeza de strings e marcadores "NULL"
    def sanitizar_texto(val):
        if pd.isna(val) or val is None:
            return None
        s = str(val).strip()
        # Elimina caracteres de controle ocultos
        s = re.sub(r'[\r\n\t\x00-\x1f\x7f-\x9f]', '', s)
        # Converte marcadores genéricos para NULL
        if s in ["???", "----", "N/A", "null", "NULL", "00/00/0000"]:
            return None
        return s

    # Aplicar higienização textual em todas as colunas de texto
    for col in df_sanitizado.columns:
        if df_sanitizado[col].dtype == 'object':
            df_sanitizado[col] = df_sanitizado[col].apply(sanitizar_texto)

    # Tratamento de Datas
    if "DATA_MEDICAO" in df_sanitizado.columns:
        def tratar_data(val):
            if val is None or pd.isna(val):
                return None
            if isinstance(val, (datetime.date, datetime.datetime)):
                return val.strftime('%Y-%m-%d')
            s = str(val).strip()
            try:
                dt = pd.to_datetime(s, dayfirst=True)
                return dt.strftime('%Y-%m-%d')
            except:
                return None # Vira NULL caso seja uma data inválida

        df_sanitizado["DATA_MEDICAO"] = df_sanitizado["DATA_MEDICAO"].apply(tratar_data)

    # Tratamento Numérico com Preservação de Sinal
    colunas_numericas = ["VALOR_OBTIDO", "VALOR_NOMINAL", "TOLERANCIA_SUP", "TOLERANCIA_INF", "DESVIO"]
    for col in colunas_numericas:
        if col in df_sanitizado.columns:
            def converter_decimal(val):
                if val is None or pd.isna(val):
                    return None
                try:
                    # Trata vírgula como ponto caso venha como string
                    s = str(val).replace(',', '.')
                    return float(s)
                except:
                    return None
            df_sanitizado[col] = df_sanitizado[col].apply(converter_decimal)

    # 3. ETAPA DE VALIDAÇÃO DE OBRIGATORIEDADE E REGRAS DE NEGÓCIO
    log_execucao.append("\n ETAPA DE VALIDAÇÃO E GESTÃO DE PENDÊNCIAS")

    campos_obrigatorios = ["COD_MMC", "PART_NO", "CARACTERISTICA", "VALOR_OBTIDO"]
    registros_validos = []

    for idx, row in df_sanitizado.iterrows():
        faltantes = []
        for campo in campos_obrigatorios:
            if campo in row and (pd.isna(row[campo]) or row[campo] is None or row[campo] == ""):
                faltantes.append(campo)

        if faltantes:
            pendencias_registradas.append({
                "LINHA_REGISTRO": idx + 1,
                "FONTE_ORIGEM": row.get("FONTE_ORIGEM", "Desconhecida"),
                "CAMPOS_FALTANTES": ", ".join(faltantes),
                "STATUS_AUDITORIA": "PENDENTE_CORRECAO"
            })
        else:
            registros_validos.append(row)

    df_validos = pd.DataFrame(registros_validos)
    df_pendencias = pd.DataFrame(pendencias_registradas)

    log_execucao.append(f"• Registros Extraídos no Total: {len(df_bruto)}")
    log_execucao.append(f"• Registros Validados para Carga: {len(df_validos)}")
    log_execucao.append(f"• Registros com Pendências Registradas: {len(df_pendencias)}")
    log_execucao.append(f"PROCESSAMENTO CONCLUÍDO COM SUCESSO!")

    texto_log = "\n".join(log_execucao)

    return texto_log, df_bruto.head(10), df_validos.head(10), df_pendencias if not df_pendencias.empty else "Nenhuma pendência encontrada."

# Interface gráfica usando Gradio


with gr.Blocks(title="Sistema ETL - Schulz S.A.") as app:
    gr.Markdown("# 🏢 Protótipo do Sistema de Ingestão e ETL - Schulz S.A.")
    gr.Markdown(
        "Demonstração interativa das regras de higienização, mapeamento de origem, "
        "preservação de códigos alfanuméricos e gestão de pendências."
    )

    with gr.Row():
        with gr.Column(scale=1):
            input_file = gr.File(label="Selecione a Planilha da Schulz (.xlsx)", file_types=[".xlsx"])
            btn_run = gr.Button("🚀 Executar Ingestão ETL", variant="primary")

        with gr.Column(scale=2):
            output_log = gr.Textbox(label="Log da Operação de ETL e Auditoria", lines=12)

    gr.Markdown("---")
    gr.Markdown("### 📊 Visualização das Etapas do Processamento")

    with gr.Tabs():
        with gr.TabItem("1. Dados Brutos (Preservando a Origem)"):
            table_raw = gr.Dataframe(interactive=False)

        with gr.TabItem("2. Base Sanitizada Pronta para Carga"):
            table_clean = gr.Dataframe(interactive=False)

        with gr.TabItem("3. Painel de Pendências (Incompletos/Auditoria)"):
            table_pendencias = gr.Dataframe(interactive=False)

    btn_run.click(
        fn=pipeline_etl_schulz,
        inputs=[input_file],
        outputs=[output_log, table_raw, table_clean, table_pendencias]
    )

# Lançar a aplicação no Colab
app.launch(share=True)
